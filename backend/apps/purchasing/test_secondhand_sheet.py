"""中古機大量進貨的 Excel(owner 2026-10-14):範本長什麼樣、讀檔讀到什麼、哪些格子會被標出來。全部是假資料。"""
import base64
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from apps.backup.tests.factory import Company
from apps.catalog.models import Category, Product
from apps.purchasing import secondhand_sheet as sheet
from apps.tenants import abilities

URL = "/api/v1/secondhand-purchase-sheet/"


def xlsx(rows, header=("品名", "IMEI", "序號", "成色", "電池", "進價", "售價", "備註"), title=sheet.DATA_SHEET):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class _Sheet(TestCase):
    def setUp(self):
        self.c = Company("a", "甲通訊行", "甲")
        used = Category.objects.create(tenant=self.c.tenant, code="SC", name="中古", is_secondhand_default=True)
        self.used = Product.objects.create(tenant=self.c.tenant, category=used, name="中古 iPhone 13 Pro 256GB 白")
        Product.objects.create(tenant=self.c.tenant, category=used, name="中古 停用的", is_active=False)
        other = Company("b", "乙通訊行", "乙")
        Product.objects.create(tenant=other.tenant, category=Category.objects.create(
            tenant=other.tenant, code="SC", name="中古", is_secondhand_default=True), name="乙的中古機")

    def upload(self, raw, client=None, name="進貨.xlsx"):
        return (client or self.c.clerk).post(URL, {"file": SimpleUploadedFile(name, raw)}, format="multipart")


class TemplateTests(_Sheet):
    def test_the_template_has_the_columns_text_codes_dropdowns_and_this_companys_used_phones(self):
        r = self.c.clerk.get(URL)
        self.assertEqual(r.status_code, 200, r.content.decode())
        self.assertEqual(r.json()["filename"], "中古機進貨範本.xlsx")
        wb = load_workbook(BytesIO(base64.b64decode(r.json()["file"])))
        self.assertEqual(wb.sheetnames, ["中古機進貨", "品名"])
        ws = wb["中古機進貨"]
        self.assertEqual([c.value for c in ws[1]], ["品名", "IMEI", "序號", "成色", "電池", "進價", "售價", "備註"])
        self.assertEqual((ws["B2"].number_format, ws["C501"].number_format), ("@", "@"))   # 長數字不會變成科學記號
        self.assertEqual(ws.max_row, 501)
        rules = {str(dv.sqref): (dv.type, dv.formula1) for dv in ws.data_validations.dataValidation}
        self.assertEqual(rules["D2:D501"], ("list", '"S,A,B,C,D"'))
        self.assertEqual((rules["F2:F501"], rules["G2:G501"]), (("whole", "1"), ("whole", "0")))
        self.assertEqual(rules["E2:E501"][0], "whole")
        self.assertEqual(rules["A2:A501"], ("list", "='品名'!$A$2:$A$2"))
        names = [tuple(c.value for c in row) for row in wb["品名"].iter_rows(min_row=2)]
        self.assertEqual(names, [("中古 iPhone 13 Pro 256GB 白", self.used.sku)])   # 停用的、別家的不列

    def test_the_template_reads_back_as_nothing_to_bring_in(self):
        raw = base64.b64decode(self.c.clerk.get(URL).json()["file"])
        r = self.upload(raw)
        self.assertEqual((r.status_code, r.json()["detail"]), (400, "沒有讀到任何一台"))


class ReadTests(_Sheet):
    def test_each_row_is_one_phone(self):
        r = self.upload(xlsx([
            ["中古 iPhone 13 Pro 256GB 白", "351135566270160", "", "b", 85, 15500, 18900, "框掉漆"],
            ["IP13P 256 白", 351135566270178, "F2LXK1AB", "a", "92%", "1,200", None, None],
            [None, None, None, None, None, None, None, None],
            ["IP12 128", "", "SN-1", "A", 0.81, 9000.4, "", ""],
        ]))
        self.assertEqual(r.status_code, 200, r.content.decode())
        rows = r.json()["rows"]
        self.assertEqual([(x["row"], x["name"], x["imei"], x["sn"], x["grade"], x["battery"], x["cost"], x["price"], x["note"],
                           x["problems"]) for x in rows], [
            (2, "中古 iPhone 13 Pro 256GB 白", "351135566270160", "", "B", 85, 15500, 18900, "框掉漆", []),
            (3, "IP13P 256 白", "351135566270178", "F2LXK1AB", "A", 92, 1200, None, "", []),  # Excel 存成數字的 IMEI 照收
            (5, "IP12 128", "", "SN-1", "A", 81, 9000, None, "", []),                         # 百分比格式 0.81、進價四捨五入
        ])

    def test_bad_cells_are_marked_on_their_row(self):
        rows = self.upload(xlsx([
            ["", "", "", "X", 120, None, -5, ""],
            ["IP13", 3.5e14 + 0.5, "", "", "abc", "免費", "", ""],
            ["IP13", "351135566270160", "", "A", "", 0, "", ""],
        ])).json()["rows"]
        self.assertEqual(rows[0]["problems"], ["品名要填", "IMEI、序號至少填一個", "成色要是 S、A、B、C、D",
                                               "電池要是 0 到 100", "進價要填", "售價不能是負的"])
        self.assertEqual(rows[1]["problems"], ["IMEI被 Excel 改成數字了", "成色要填", "電池要是數字", "進價要是數字"])
        # 複審 P0:進價 0 存檔時會被當成沒填、改用那一行的單價 → 不收
        self.assertEqual(rows[2]["problems"], ["進價要大於 0"])

    def test_columns_can_be_in_any_order_and_a_sheet_with_another_name_is_read(self):
        rows = self.upload(xlsx([["351135566270160", 500, "IP13", "B"]], header=("IMEI", "進價", "品名", "成色"), title="工作表1")).json()["rows"]
        self.assertEqual((rows[0]["name"], rows[0]["imei"], rows[0]["cost"]), ("IP13", "351135566270160", 500))

    def test_a_file_it_cannot_read_is_refused(self):
        for raw, says in [
            (b"not a zip", "要是 .xlsx 檔"),
            (xlsx([["IP13", "1"]], header=("品名", "IMEI")), "第一列要有「品名」「IMEI」或「序號」「進價」"),
            (b"PK\x03\x04broken", "讀不出這個檔"),
        ]:
            r = self.upload(raw)
            self.assertEqual((r.status_code, r.json()["detail"]), (400, says))
        self.assertEqual(self.c.clerk.post(URL, {}, format="multipart").json()["detail"], "要附上 Excel 檔")

    def test_too_many_rows_are_refused_not_cut_short(self):
        rows = [["IP13", f"35113556627{i:04d}", "", "A", "", 1, "", ""] for i in range(sheet.MAX_ROWS + 1)]
        r = self.upload(xlsx(rows))
        self.assertEqual((r.status_code, r.json()["detail"]), (400, f"一次最多 {sheet.MAX_ROWS} 台"))

    def test_needs_the_purchase_item(self):
        from apps.tenants.models import UserProfile

        UserProfile.objects.filter(user=self.c.clerk_user).update(denied_abilities=[abilities.PURCHASE])
        clerk = Company.client(get_user_model().objects.get(pk=self.c.clerk_user.pk))
        self.assertEqual(clerk.get(URL).status_code, 403)
        self.assertEqual(self.upload(xlsx([["IP13", "351135566270160", "", "A", "", 1, "", ""]]), clerk).status_code, 403)


class CheckTests(_Sheet):
    """預覽:品名只認確定的(紅隊第 1 條)、碼已經在系統裡的標出來(紅隊第 2 條)。"""

    def rows(self, data, **extra):
        r = self.c.clerk.post(URL, {"file": SimpleUploadedFile("進貨.xlsx", xlsx(data)), **extra}, format="multipart")
        self.assertEqual(r.status_code, 200, r.content.decode())
        return r.json()

    def test_only_a_sure_name_is_matched(self):
        from apps.identity.models import ProductAlias

        ProductAlias.objects.create(tenant=self.c.tenant, product=self.used, kind=ProductAlias.Kind.LEGACY_NAME,
                                    value="IP13P 256 白", verified=True)
        out = self.rows([
            ["中古 iPhone 13 Pro 256GB 白", "351135566270160", "", "A", "", 1, "", ""],
            [self.used.sku, "351135566270178", "", "A", "", 1, "", ""],
            ["IP13P 256 白", "351135566270186", "", "A", "", 1, "", ""],             # 記住過的叫法
            ["中古 iPhone 13 Pro 256GB", "351135566270194", "", "A", "", 1, "", ""],  # 很像,但不確定 → 不帶
            ["中古 iPhone 13 Pro 256GB", "351135566270202", "", "A", "", 1, "", ""],
            ["中古 停用的", "351135566270210", "", "A", "", 1, "", ""],
            ["乙的中古機", "351135566270228", "", "A", "", 1, "", ""],
        ])
        self.assertEqual([r["product"] for r in out["rows"]], [self.used.id] * 3 + [None] * 4)
        self.assertEqual(out["names"], [{"name": "中古 iPhone 13 Pro 256GB", "count": 2},
                                        {"name": "中古 停用的", "count": 1}, {"name": "乙的中古機", "count": 1}])
        [p] = out["products"]
        self.assertEqual((p["id"], p["is_secondhand"], p["requires_serial"]), (self.used.id, True, True))

    def test_a_suppliers_own_name_only_counts_for_that_supplier(self):
        from apps.identity.models import ProductAlias
        from apps.parties.models import Supplier

        other = Supplier.objects.create(tenant=self.c.tenant, name="甲另一家")
        ProductAlias.objects.create(tenant=self.c.tenant, product=self.used, kind=ProductAlias.Kind.VENDOR_NAME,
                                    supplier=self.c.supplier, value="13P-256-W", verified=True)
        row = ["13P-256-W", "351135566270160", "", "A", "", 1, "", ""]
        self.assertEqual(self.rows([row], supplier=self.c.supplier.id)["rows"][0]["product"], self.used.id)
        self.assertIsNone(self.rows([row], supplier=other.id)["rows"][0]["product"])
        r = self.c.clerk.post(URL, {"file": SimpleUploadedFile("x.xlsx", xlsx([row])), "supplier": 99999}, format="multipart")
        self.assertEqual((r.status_code, r.json()["detail"]), (400, "找不到這個廠商"))

    def test_codes_already_in_the_system_or_twice_in_the_file_are_marked(self):
        self.c.purchase(phone_serials=["351135566270160", "351135566270178"])
        self.c.sell(serial_no="351135566270178")
        out = self.rows([
            ["中古 iPhone 13 Pro 256GB 白", "351135566270160", "", "A", "", 1, "", ""],
            ["中古 iPhone 13 Pro 256GB 白", "351135566270178", "", "A", "", 1, "", ""],
            ["中古 iPhone 13 Pro 256GB 白", "351135566270186", "", "A", "", 1, "", ""],
            ["中古 iPhone 13 Pro 256GB 白", "351-1355-6627-0186", "", "A", "", 1, "", ""],
            ["中古 iPhone 13 Pro 256GB 白", "", "sn-9", "A", "", 1, "", ""],
            ["中古 iPhone 13 Pro 256GB 白", "", "SN9", "A", "", 1, "", ""],
        ])
        self.assertEqual([r["problems"] for r in out["rows"]], [
            [f"351135566270160 已經在系統裡({self.c.wh.name} 在庫)"],
            ["351135566270178 以前賣出去過,還不能從這裡收回來"],
            [],
            ["351-1355-6627-0186 跟第 4 列重複"],
            [],
            ["SN9 跟第 6 列重複"],
        ])
