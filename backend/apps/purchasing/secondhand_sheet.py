"""中古機大量進貨的 Excel(owner 2026-10-14:「中古機我希望可以用Excel匯入的方式大量進貨，所以幫我設計好大量匯入的Excel範本」)。

- 範本(`build_template`):一列一台 —— 品名、IMEI、序號、成色、電池、進價、售價、備註。廠商、門市、日期、課稅別在畫面上選,不放 Excel。
  附一張「品名」= 這家公司現有的中古機品號,品名那一格可以從清單選、也可以直接貼廠商的寫法(不擋)。
  IMEI / 序號那兩欄預先設成文字(不然 Excel 會把長數字變成科學記號);成色下拉;電池、價錢只收整數。
- 讀檔(`read_sheet`):只讀、不寫任何東西;回每一列讀到的與每一格的問題。品名對到哪個品號、碼在不在系統裡,
  在畫面上用原本的比對(`products/resolve/`)與序號檢查(`serials/check/`)做 —— 跟手動刷、批次貼上同一套。
- 加進進貨單之後跟一台一台刷的一樣:店員按儲存才成立,庫存、成本、序號的規則全部照進貨那一份(`services.commit_purchase_order`)。
- **成色、進價一定要有、進價要大於 0**(複審 P0):進貨單上中古機一定要選成色;每一台的成本是 0 的話存檔時會改用那一行的單價
  (`services._serial_cost` 把 0 當成沒填),匯入的那一行單價是 0、併進去的那一行是別的價 —— 成本就不是這一列寫的。
"""
from __future__ import annotations

from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

DATA_SHEET = "中古機進貨"
NAMES_SHEET = "品名"
# (欄名, 寬度);欄名就是讀檔時認的名字(順序可以換)
COLUMNS = [("品名", 34), ("IMEI", 18), ("序號", 16), ("成色", 6), ("電池", 6), ("進價", 10), ("售價", 10), ("備註", 30)]
TEXT_COLUMNS = ("IMEI", "序號")
GRADES = ("S", "A", "B", "C", "D")
TEMPLATE_ROWS = 500
MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024
MAX_PRICE = 100_000_000


class SheetError(Exception):
    """整份看不懂(原因可以直接給人看)。"""


def build_template(tenant) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    from apps.catalog.models import Product

    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    names = wb.create_sheet(NAMES_SHEET)
    names.append(["品名", "品號"])
    rows = (Product.objects.filter(tenant=tenant, is_secondhand=True, is_active=True)
            .order_by("name").values_list("name", "sku"))
    for name, sku in rows:
        names.append([name, sku])
    names.column_dimensions["A"].width = 40
    names.column_dimensions["B"].width = 14
    names.freeze_panes = "A2"

    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDE3EA")
    for i, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=title)
        cell.font, cell.fill = bold, fill
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if title in TEXT_COLUMNS:
            for r in range(2, TEMPLATE_ROWS + 2):
                ws.cell(row=r, column=i).number_format = "@"
    ws.freeze_panes = "A2"
    col = {title: get_column_letter(i) for i, (title, _) in enumerate(COLUMNS, start=1)}
    last = TEMPLATE_ROWS + 1

    def rng(title):
        return f"{col[title]}2:{col[title]}{last}"

    checks = []
    if rows:
        # 從清單選;不在清單上的照收(貼廠商的寫法),到畫面上再對
        pick = DataValidation(type="list", formula1=f"='{NAMES_SHEET}'!$A$2:$A${len(rows) + 1}", allow_blank=True)
        pick.showErrorMessage = False
        pick.add(rng("品名"))
        checks.append(pick)
    grade = DataValidation(type="list", formula1='"' + ",".join(GRADES) + '"', allow_blank=True)
    grade.add(rng("成色"))
    battery = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
    battery.add(rng("電池"))
    cost = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True)
    cost.add(rng("進價"))
    money = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    money.add(rng("售價"))
    for dv in (*checks, grade, battery, cost, money):
        ws.add_data_validation(dv)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _code(value, title, problems) -> str:
    """IMEI / 序號:要是文字。Excel 存成數字的整數照收(15 碼以內存得完整);有小數的 = 被改過,擋下。"""
    if isinstance(value, float) and not value.is_integer():
        problems.append(f"{title}被 Excel 改成數字了")
        return ""
    return _text(value).replace(" ", "")


def _whole(value, title, problems, *, required=False, top=MAX_PRICE, percent=False, positive=False) -> int | None:
    text = _text(value).replace(",", "").rstrip("%").strip()
    if text == "":
        if required:
            problems.append(f"{title}要填")
        return None
    try:
        number = Decimal(text)
    except InvalidOperation:
        problems.append(f"{title}要是數字")
        return None
    if percent and isinstance(value, float) and 0 < number <= 1:
        number *= 100                      # 設成百分比格式的格子讀出來是 0.85
    number = number.quantize(Decimal(1), rounding=ROUND_HALF_UP)
    if percent and not 0 <= number <= top:
        problems.append(f"{title}要是 0 到 {top}")
        return None
    if number < 0:
        problems.append(f"{title}不能是負的")
        return None
    if positive and number == 0:
        problems.append(f"{title}要大於 0")
        return None
    if number > top:
        problems.append(f"{title}太大")
        return None
    return int(number)


def read_sheet(raw: bytes) -> list[dict]:
    """範本(或照範本欄名的任何 xlsx)→ 每一台一列:{row, name, imei, sn, grade, battery, cost, price, note, problems}。"""
    from openpyxl import load_workbook

    if len(raw) > MAX_BYTES:
        raise SheetError("檔案太大(最多 5 MB)")
    if raw[:2] != b"PK":
        raise SheetError("要是 .xlsx 檔")
    try:
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise SheetError("讀不出這個檔") from None
    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [_text(h) for h in next(rows, ())]
    col = {h: i for i, h in enumerate(header) if h}
    if "品名" not in col or ("IMEI" not in col and "序號" not in col) or "進價" not in col:
        raise SheetError("第一列要有「品名」「IMEI」或「序號」「進價」")

    def cell(values, title):
        i = col.get(title)
        return values[i] if i is not None and i < len(values) else None

    out = []
    for n, values in enumerate(rows, start=2):
        if not any(_text(v) for v in values):
            continue
        if len(out) >= MAX_ROWS:
            raise SheetError(f"一次最多 {MAX_ROWS} 台")
        problems: list[str] = []
        name = _text(cell(values, "品名"))
        if not name:
            problems.append("品名要填")
        imei, sn = _code(cell(values, "IMEI"), "IMEI", problems), _code(cell(values, "序號"), "序號", problems)
        if not imei and not sn and not any("被 Excel" in p for p in problems):
            problems.append("IMEI、序號至少填一個")
        grade = _text(cell(values, "成色")).upper()
        if not grade:
            problems.append("成色要填")
        elif grade not in GRADES:
            problems.append("成色要是 S、A、B、C、D")
            grade = ""
        out.append({
            "row": n, "name": name[:200], "imei": imei[:80], "sn": sn[:80], "grade": grade,
            "battery": _whole(cell(values, "電池"), "電池", problems, top=100, percent=True),
            "cost": _whole(cell(values, "進價"), "進價", problems, required=True, positive=True),
            "price": _whole(cell(values, "售價"), "售價", problems),
            "note": _text(cell(values, "備註"))[:200],
            "problems": problems,
        })
    if not out:
        raise SheetError("沒有讀到任何一台")
    return out


def match_name(tenant, name, supplier=None, _cache=None):
    """品名 → 這家公司啟用中的中古機品號。**只認確定的**(紅隊:像但不確定的不帶,店員會整批按下去):
    一字不差的品名、品號、條碼,或之前記住過、已確認的叫法(通用的,或這家廠商的)。其餘回 None,要人自己挑。"""
    from apps.catalog.models import Product
    from apps.identity.product_match import identifier_hits

    name = (name or "").strip()
    if not name:
        return None
    usable = Product.objects.filter(tenant=tenant, is_secondhand=True, is_active=True)
    same = list(usable.filter(name=name).values_list("pk", flat=True)[:2])
    if len(same) == 1:
        return same[0]
    ids = set(identifier_hits(tenant, name, supplier=supplier)) & set(usable.values_list("pk", flat=True))
    return next(iter(ids)) if len(ids) == 1 else None


def check_rows(tenant, rows, supplier=None) -> dict:
    """讀出來的每一列加上:對到哪個品號(`product`:編號或 None)、碼的問題(這份檔裡重複、已經在系統裡)。
    回 {rows, names}:names = 對不到的品名與幾台(畫面上一個品名挑一次)。只讀。"""
    from apps.identity.normalize import normalize_serial
    from apps.inventory.identifiers import taken_details
    from apps.inventory.models import ProductSerial

    S = ProductSerial.Status
    found = {}
    for row in rows:
        if row["name"] and row["name"] not in found:
            found[row["name"]] = match_name(tenant, row["name"], supplier)
    seen: dict[str, int] = {}
    owners = taken_details(tenant, [c for row in rows for c in (row["imei"], row["sn"]) if c])
    for row in rows:
        row["product"] = found.get(row["name"])
        for code in (row["imei"], row["sn"]):
            if not code:
                continue
            key = normalize_serial(code)
            if key in seen:
                row["problems"].append(f"{code} 跟第 {seen[key]} 列重複")
                continue
            seen[key] = row["row"]
            serial = owners.get(code)
            if serial is None:
                continue
            if serial.status == S.SOLD:
                # 賣出去的手機收回來另外排(owner 2026-10-10);現在入不了
                row["problems"].append(f"{code} 以前賣出去過,還不能從這裡收回來")
            else:
                where = " ".join(x for x in (serial.warehouse.name if serial.warehouse_id else "",
                                             serial.get_status_display()) if x)
                row["problems"].append(f"{code} 已經在系統裡({where})")
    names: dict[str, int] = {}
    for row in rows:
        if row["name"] and row["product"] is None:
            names[row["name"]] = names.get(row["name"], 0) + 1
    return {"rows": rows, "names": [{"name": k, "count": v} for k, v in names.items()]}
