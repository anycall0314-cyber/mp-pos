"""舊 POS 有庫存的品項整批建成品號(12/1 盤點入庫之前;owner 2026-10-14)。全部是假資料。"""
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
import shutil
import tempfile

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook

from apps.catalog.models import Category, Product
from apps.identity.models import ProductAlias
from apps.inventory.models import StockBalance, Warehouse
from apps.legacy import products
from apps.legacy.archive import SOURCE_SYSTEM
from apps.legacy.models import LegacyMappingLog, LegacyProductMap, MapMethod, MapStatus
from apps.tenants.models import Tenant

HEAD_STOCK = ["序號", "分類", "品號", "品名", "店名", "數量", "調出中(未確認)", "待銷中(未確認)", "可用量", "進貨提示(安全存量-數量)"]
HEAD_PRICE = ["序號", "類別", "品號", "品名", "共用", "零售價", "民生店", "湳雅店", "維修倉", "小計"]
HEAD_COST = ["序號", "品號", "品名", "民生店", "湳雅店", "成本", "總數量", "總成本"]


def ored_page(header, rows, repeat_every=2):
    """歐睿匯出的樣子:Big5 網頁表格、前面一段標題與程式、表頭每幾列重複一次、最後一列小計。"""
    out = ["<html><head><script>var x = '<tr>';</script></head><body><table>",
           "<tr><td>庫 存 明 細 表</td><td>製表時間:2026/10/7</td></tr>"]
    for i, r in enumerate(rows):
        if i % repeat_every == 0:
            out.append("<tr>" + "".join(f"<td>{h}</td>" for h in header) + "</tr>")
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in [str(i + 1), *r]) + "</tr>")
    out.append("<tr><td>小計</td><td></td></tr></table></body></html>")
    return "\n".join(out).encode("big5hkscs")


def price_rows(*rows):
    """(品號, 品名, 零售價, 小計) → 商品價格查詢表的一列。"""
    return [["T12", code, name, "", price, "0", total, "0", total] for code, name, price, total in rows]


MAPPING = {"皮套": "皮套", "9H": "玻璃貼", "充電": "充電/線材", "中古": "中古機", "維修料": "維修料",
           "__excluded__": {"遠傳", "手機"}}


class ReadReportTests(TestCase):
    def test_stock_detail_report_adds_up_every_store(self):
        raw = ored_page(HEAD_STOCK, [
            ["T12", "T12SAMNOTE502", "皮套/SAM-NOTE5-黑", "民生店", "2", "0", "0", "2", ""],
            ["T12", "T12SAMNOTE502", "皮套/SAM-NOTE5-黑", "湳雅店", "3", "0", "0", "3", ""],
            ["T12", "T12SAMNOTE502", "皮套/SAM-NOTE5-黑", "亞太巨城", "1", "0", "0", "1", ""],
            ["T20", "  T20ASUS005  ", "9H/ASUS/ZFC", "民生店", "0", "0", "0", "0", ""],
        ])
        items = {i.code: i for i in products.stock_items(products.read_ored_rows(raw))}
        self.assertEqual(sorted(items), ["T12SAMNOTE502", "T20ASUS005"])     # 頭尾空白不算品號的一部分;表頭重複、小計列都跳過
        self.assertEqual((items["T12SAMNOTE502"].qty, items["T12SAMNOTE502"].price), (6, None))   # 每一家店都算;這一張沒有售價
        self.assertEqual(items["T20ASUS005"].qty, 0)

    def test_price_report_reads_the_retail_price_as_whole_dollars(self):
        raw = ored_page(HEAD_PRICE, price_rows(
            ("T4710000000001", "9H/IMOS/IP15PM/防窺", "1,290", "3"),
            ("T25005", "充電/卡準式線套", "99.5", "1"),
            ("T25006", "充電/I線套", "", "2"),
        ))
        items = {i.code: i for i in products.stock_items(products.read_ored_rows(raw))}
        self.assertEqual([(items[c].price, items[c].qty) for c in ("T4710000000001", "T25005", "T25006")],
                         [(1290, 3), (100, 1), (None, 2)])                     # 99.5 四捨五入;空的 = 沒有售價

    def test_in_stock_means_any_store_above_zero_not_the_total(self):
        raw = ored_page(HEAD_PRICE, [["T12", "T12A", "皮套/A", "", "390", "2", "-2", "0", "0"],
                                     ["T12", "T12B", "皮套/B", "", "390", "0", "0", "0", "0"]])
        items = {i.code: i for i in products.stock_items(products.read_ored_rows(raw))}
        self.assertEqual((items["T12A"].qty, items["T12A"].in_stock, items["T12B"].in_stock), (0, True, False))
        stock = ored_page(HEAD_STOCK, [["T12", "T12A", "皮套/A", "民生店", "2", "0", "0", "2", ""],
                                       ["T12", "T12A", "皮套/A", "湳雅店", "-2", "0", "0", "-2", ""]])
        [item] = products.stock_items(products.read_ored_rows(stock))
        self.assertEqual((item.qty, item.in_stock), (0, True))

    def test_a_real_xlsx_is_read_too(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["商品價格查詢表"])
        ws.append(HEAD_PRICE)
        ws.append([1, "T12", "T12A", "皮套/A", "", 350, 1, 0, 0, 1])
        buf = BytesIO()
        wb.save(buf)
        [item] = products.stock_items(products.read_ored_rows(buf.getvalue()))
        self.assertEqual((item.code, item.name, item.price, item.qty), ("T12A", "皮套/A", 350, 1))

    def test_a_report_it_cannot_read_is_refused(self):
        with self.assertRaises(products.ReportError):
            products.read_ored_rows("<table><tr><td>甲</td><td>乙</td></tr></table>".encode("big5"))
        with self.assertRaises(products.ReportError):
            products.stock_items([{"品號": "A", "品名": "皮套/A"}])           # 沒有數量 / 小計,分不出有沒有庫存

    def test_cost_report_gives_each_old_code_its_cost(self):
        raw = ored_page(HEAD_COST, [["0004A176001W3", "手機/三星-A17", "0", "2", "7,050", "2", "14100"],
                                    ["T12A", "皮套/A", "1", "0", "95.24", "1", "95"],
                                    ["T12B", "皮套/B", "1", "0", "-5", "1", "-5"],
                                    ["T12C", "皮套/C", "1", "0", "", "1", ""],
                                    ["T12D", "皮套/D", "1", "0", "10.005", "1", "10"]])
        costs = products.old_costs(products.read_ored_rows(raw))
        self.assertEqual(costs, {"0004A176001W3": Decimal("7050"), "T12A": Decimal("95.24"),     # 成本不收整數;負的、空的不收
                                 "T12D": Decimal("10.01")})                                     # 存到分
        with self.assertRaises(products.ReportError):
            products.old_costs([{"品號": "A", "品名": "皮套/A"}])

    def test_category_mapping_uses_the_owner_edit_and_keeps_exclusions_apart(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "類別對照"
        ws.append(["提議類別", "商品型態", "舊系統品名前綴", "庫存筆數", "你的修改(類別/排除)"])
        ws.append(["保護殼/背蓋", "配件", "空壓", 254, "空壓殼"])
        ws.append(["保護殼/背蓋", "配件", "邊框", 21, None])
        ws.append(["其他配件", "配件", "庫存", 112, "排除"])
        ws.append(["(排除)", "—", "遠傳", 11, "(排除)"])
        buf = BytesIO()
        wb.save(buf)
        m = products.load_category_mapping(buf.getvalue())
        self.assertEqual((m["空壓"], m["邊框"]), ("空壓殼", "保護殼/背蓋"))
        self.assertEqual(m["__excluded__"], {"庫存", "遠傳"})
        self.assertNotIn("庫存", m)

    def test_skip_list_reads_the_old_codes_on_the_delete_sheet(self):
        wb = Workbook()
        keep = wb.active
        keep.title = "舊品號"
        keep.append(["舊品號", "品名"])
        keep.append(["T12KEEP", "皮套/留著"])
        drop = wb.create_sheet("未來會刪除")
        drop.append(["舊品號", "品名", "類別"])
        drop.append(["  T12A ", "皮套/A", "皮套"])
        drop.append([None, "", ""])
        drop.append(["T20B", "9H/B", "玻璃貼"])
        buf = BytesIO()
        wb.save(buf)
        self.assertEqual(products.load_skip_list(buf.getvalue()), {"T12A", "T20B"})   # 只讀刪除那一頁;空白列不算
        other = Workbook()
        other.active.append(["品號"])
        other.active.append(["X1"])
        buf = BytesIO()
        other.save(buf)
        self.assertEqual(products.load_skip_list(buf.getvalue()), {"X1"})            # 沒有那一頁:讀第一頁、「品號」也認
        bad = Workbook()
        bad.active.append(["品名"])
        buf = BytesIO()
        bad.save(buf)
        with self.assertRaises(products.ReportError):
            products.load_skip_list(buf.getvalue())
        with self.assertRaises(products.ReportError):
            products.load_skip_list(b"<table></table>")


class BringInTests(TestCase):
    def setUp(self):
        self.a = Tenant.objects.create(name="甲通訊行", code="a")
        self.b = Tenant.objects.create(name="乙通訊行", code="b")
        self.case_cat = Category.objects.create(tenant=self.a, code="LC", name="皮套")

    def item(self, code, name, price=None, qty=1):
        return products.LegacyItem(code, name, price, qty)

    def run_it(self, items, write=True, tenant=None):
        return products.bring_in_products(tenant or self.a, items, MAPPING, write=write)

    def test_each_old_item_in_stock_becomes_one_product_without_stock_or_cost(self):
        r = self.run_it([
            self.item("T12SAMNOTE502", "皮套/SAM-NOTE5-黑", 390, qty=3),
            self.item("T20ASUS005", "9H/ASUS/ZFC", 0),
            self.item("T25006", "充電/I線套", None, qty=0),             # 沒有庫存
        ])
        self.assertEqual((r.created, r.zero_stock, r.no_price), (2, 1, 1))
        case = Product.objects.get(tenant=self.a, barcode="T12SAMNOTE502")
        self.assertEqual((case.name, case.category_id, int(case.list_price)), ("皮套/SAM-NOTE5-黑", self.case_cat.pk, 390))
        self.assertTrue(case.sku.startswith("LC-"))                        # 品號由系統編,舊品號放在條碼
        self.assertEqual(case.lifecycle_status, Product.LifecycleStatus.PENDING)
        self.assertFalse(case.requires_serial)
        self.assertEqual(int(case.weighted_avg_cost), 0)
        self.assertFalse(StockBalance.objects.filter(tenant=self.a).exists())   # 數量是盤點時刷的
        glass = Product.objects.get(tenant=self.a, barcode="T20ASUS005")
        self.assertEqual(glass.category.name, "玻璃貼")                     # 對照表的類別,沒有就建
        self.assertEqual(r.new_categories, [f"{glass.category.code} 玻璃貼"])
        self.assertEqual(glass.category.code, "PP")
        maps = LegacyProductMap.objects.filter(tenant=self.a, source_system=SOURCE_SYSTEM)
        self.assertEqual({(m.product_code_raw, m.product_id, m.status, m.method) for m in maps}, {
            ("T12SAMNOTE502", case.pk, MapStatus.CONFIRMED, MapMethod.CREATED),
            ("T20ASUS005", glass.pk, MapStatus.CONFIRMED, MapMethod.CREATED),
        })
        self.assertEqual(LegacyMappingLog.objects.filter(tenant=self.a, note=products.BULK_NOTE).count(), 2)

    def test_old_cost_becomes_the_cost_a_stocktake_brings_it_in_at(self):
        items = products.with_costs([self.item("T12A", "皮套/A", 390), self.item("T12B", "皮套/B", 390)],
                                    {"T12A": Decimal("95.24")})
        r = self.run_it(items)
        self.assertEqual(r.no_cost, 1)
        self.assertEqual(Product.objects.get(tenant=self.a, barcode="T12A").weighted_avg_cost, Decimal("95.24"))
        self.assertEqual(Product.objects.get(tenant=self.a, barcode="T12B").weighted_avg_cost, 0)

    def test_a_stocktake_brings_the_new_product_in_at_the_old_cost(self):
        from apps.inventory import stockcount
        from django.contrib.auth import get_user_model
        from apps.tenants.models import UserProfile

        boss = get_user_model().objects.create_user("a-boss", password="pw-12345")
        UserProfile.objects.create(user=boss, tenant=self.a, role="tenant_admin", is_warehouse_locked=False)
        self.run_it(products.with_costs([self.item("T12A", "皮套/A", 390)], {"T12A": Decimal("95.24")}))
        p = Product.objects.get(tenant=self.a, barcode="T12A")
        wh = Warehouse.objects.create(tenant=self.a, code="A02", name="湳雅店")
        count = stockcount.open_count(self.a, wh, [], boss)
        stockcount.set_line(count, p, 4, boss)
        stockcount.submit(count, boss)
        stockcount.approve(count, boss)
        balance = StockBalance.objects.get(tenant=self.a, product=p, warehouse=wh)
        self.assertEqual((balance.qty, balance.weighted_avg_cost), (4, Decimal("95.24")))

    def test_an_existing_product_keeps_its_own_cost(self):
        mine = Product.objects.create(tenant=self.a, category=self.case_cat, name="皮套 A", barcode="T12A",
                                      weighted_avg_cost=Decimal("80"))
        self.run_it(products.with_costs([self.item("T12A", "皮套/A", 390)], {"T12A": Decimal("95")}))
        self.assertEqual(Product.objects.get(pk=mine.pk).weighted_avg_cost, Decimal("80"))

    def test_categories_unrelated_to_a_phone_model_do_not_ask_for_one(self):
        self.run_it([self.item("C1", "充電/線", 100), self.item("G1", "9H/IP15", 100)])
        self.assertFalse(Category.objects.get(tenant=self.a, name="充電/線材").needs_host_model)
        self.assertTrue(Category.objects.get(tenant=self.a, name="玻璃貼").needs_host_model)

    def test_repair_parts_go_to_the_parts_store(self):
        self.run_it([self.item("R1", "維修料/IP11/螢幕", 1500), self.item("C1", "充電/線", 100)])
        self.assertEqual(Product.objects.get(tenant=self.a, barcode="R1").warehouse_type, Product.WarehouseType.PARTS)
        self.assertEqual(Product.objects.get(tenant=self.a, barcode="C1").warehouse_type, Product.WarehouseType.PRODUCT)

    def test_phones_excluded_and_unmapped_items_are_not_created_but_counted(self):
        r = self.run_it([
            self.item("356000000000001I", "中古/IP13/128G/藍/A"),
            self.item("0001XT2055404L", "手機/MOTO/XT2055"),
            self.item("AA01", "遠傳/月租"),
            self.item("F1", "富力森輕巧頸式風扇-白"),
            self.item("F2", "福利機/VIVO/V60"),
        ])
        self.assertEqual((r.created, r.serial), (0, 1))
        self.assertEqual(r.excluded, {"手機": 1, "遠傳": 1})
        self.assertEqual(r.unmapped, {"富力森輕巧頸式風扇-白": 1, "福利機": 1})
        self.assertFalse(Product.objects.filter(tenant=self.a).exists())

    def test_the_trial_run_writes_nothing_and_gives_the_same_numbers(self):
        items = [self.item("T12A", "皮套/A", 100), self.item("T4710000000001", "9H/B", 200)]
        trial = self.run_it(items, write=False)
        self.assertFalse(Product.objects.filter(tenant=self.a).exists())
        self.assertFalse(Category.objects.filter(tenant=self.a, name="玻璃貼").exists())
        self.assertFalse(LegacyProductMap.objects.exists())
        self.assertFalse(ProductAlias.objects.exists())
        real = self.run_it(items)
        self.assertEqual(trial, real)

    def test_run_again_leaves_what_was_done(self):
        items = [self.item("T12A", "皮套/A", 100)]
        self.run_it(items)
        again = self.run_it(items + [self.item("T12B", "皮套/B", 120)])
        self.assertEqual((again.created, again.already), (1, 1))
        self.assertEqual(Product.objects.filter(tenant=self.a).count(), 2)

    def test_t_plus_barcode_also_scans_by_the_maker_barcode(self):
        self.run_it([self.item("T4710000000001", "皮套/A", 100)])
        p = Product.objects.get(tenant=self.a, barcode="T4710000000001")
        alias = ProductAlias.objects.get(tenant=self.a, product=p)
        self.assertEqual((alias.kind, alias.value, alias.verified, alias.source),
                         (ProductAlias.Kind.BARCODE, "4710000000001", True, ProductAlias.Source.IMPORTED))

    def test_maker_barcode_already_on_another_product_is_left_alone(self):
        other = Product.objects.create(tenant=self.a, category=self.case_cat, name="原廠皮套", barcode="4710000000001")
        r = self.run_it([self.item("T4710000000001", "皮套/A", 100)])
        self.assertEqual((r.created, r.aliases), (1, 0))
        self.assertFalse(ProductAlias.objects.filter(tenant=self.a).exists())
        self.assertEqual(Product.objects.get(pk=other.pk).barcode, "4710000000001")

    def test_a_product_with_this_barcode_is_linked_not_duplicated(self):
        mine = Product.objects.create(tenant=self.a, category=self.case_cat, name="皮套 NOTE5 黑", barcode="T12SAMNOTE502")
        via_alias = Product.objects.create(tenant=self.a, category=self.case_cat, name="皮套 S7")
        ProductAlias.objects.create(tenant=self.a, product=via_alias, kind=ProductAlias.Kind.BARCODE, value="T12S7")
        # 別家公司同一個條碼不算
        Product.objects.create(tenant=self.b, category=Category.objects.create(tenant=self.b, code="LC", name="皮套"),
                               name="乙的皮套", barcode="T12C")
        r = self.run_it([self.item("T12SAMNOTE502", "皮套/SAM-NOTE5-黑", 390),
                         self.item("T12S7", "皮套/S7", 390), self.item("T12C", "皮套/C", 100)])
        self.assertEqual((r.created, r.linked), (1, 2))
        self.assertEqual(Product.objects.get(pk=mine.pk).name, "皮套 NOTE5 黑")      # 原本有的不改
        maps = {m.product_code_raw: m for m in LegacyProductMap.objects.filter(tenant=self.a)}
        self.assertEqual((maps["T12SAMNOTE502"].product_id, maps["T12SAMNOTE502"].method), (mine.pk, MapMethod.CODE))
        self.assertEqual(maps["T12S7"].product_id, via_alias.pk)
        self.assertEqual(Product.objects.get(tenant=self.a, barcode="T12C").name, "皮套/C")

    def test_same_name_with_another_barcode_is_listed_not_created(self):
        Product.objects.create(tenant=self.a, category=self.case_cat, name="皮套/A", barcode="X1")
        r = self.run_it([self.item("T12A", "皮套/A", 100)])
        self.assertEqual((r.created, r.same_name), (0, [("T12A", "皮套/A")]))

    def test_an_old_code_seen_in_history_gets_confirmed(self):
        seen = LegacyProductMap.objects.create(tenant=self.a, source_system=SOURCE_SYSTEM, product_code_raw="T12A",
                                               product_name_seen="皮套/A(舊名)")
        self.run_it([self.item("T12A", "皮套/A", 100)])
        seen.refresh_from_db()
        self.assertEqual((seen.status, seen.product.name, seen.product_name_seen),
                         (MapStatus.CONFIRMED, "皮套/A", "皮套/A(舊名)"))
        self.assertEqual(LegacyProductMap.objects.filter(tenant=self.a).count(), 1)

    def test_old_codes_the_owner_picked_are_not_created_or_linked(self):
        mine = Product.objects.create(tenant=self.a, category=self.case_cat, name="皮套 S7", barcode="T12S7")
        r = products.bring_in_products(self.a, [
            self.item("T12A", "皮套/A", 100), self.item("T12B", "皮套/B", 100),
            self.item("T12S7", "皮套/S7", 100), self.item("T12Z", "皮套/Z", 100, qty=0),
        ], MAPPING, write=True, skip={"T12A", "T12S7", "T12Z"})
        self.assertEqual((r.created, r.skipped, r.linked, r.zero_stock), (1, 3, 0, 0))
        self.assertEqual(set(Product.objects.filter(tenant=self.a).values_list("barcode", flat=True)), {"T12B", "T12S7"})
        self.assertEqual(list(LegacyProductMap.objects.filter(tenant=self.a).values_list("product_code_raw", flat=True)), ["T12B"])
        self.assertTrue(Product.objects.filter(pk=mine.pk).exists())                      # 原本就有的商品不動

    def test_a_category_code_already_taken_gets_another(self):
        Category.objects.create(tenant=self.a, code="PP", name="包膜")
        self.run_it([self.item("G1", "9H/IP15", 100)])
        self.assertEqual(Category.objects.get(tenant=self.a, name="玻璃貼").code, "G01")


class TakeBackTests(TestCase):
    def setUp(self):
        self.a = Tenant.objects.create(name="甲通訊行", code="a")
        self.cat = Category.objects.create(tenant=self.a, code="LC", name="皮套")
        self.existing = Product.objects.create(tenant=self.a, category=self.cat, name="皮套 原本有", barcode="T12OLD")
        products.bring_in_products(self.a, [
            products.LegacyItem("T12A", "皮套/A", 100, 1),
            products.LegacyItem("T4710000000001", "皮套/B", 100, 1),
            products.LegacyItem("T12OLD", "皮套/原本有", 100, 1),
        ], MAPPING, write=True)
        self.used = Product.objects.get(tenant=self.a, barcode="T12A")
        wh = Warehouse.objects.create(tenant=self.a, code="A01", name="湳雅店")
        StockBalance.objects.create(tenant=self.a, product=self.used, warehouse=wh, qty=3)   # 盤點入了庫 = 用過

    def test_removes_unused_created_keeps_used_and_only_unlinks_existing(self):
        trial = products.take_back_products(self.a)
        self.assertEqual(trial, {"removed": 1, "unlinked": 1, "kept": 1})
        self.assertEqual(Product.objects.filter(tenant=self.a).count(), 3)                # 試算不刪
        r = products.take_back_products(self.a, write=True)
        self.assertEqual(r, trial)
        left = set(Product.objects.filter(tenant=self.a).values_list("barcode", flat=True))
        self.assertEqual(left, {"T12A", "T12OLD"})
        self.assertFalse(ProductAlias.objects.filter(tenant=self.a).exists())             # 原廠條碼跟著那個品號一起走
        maps = {m.product_code_raw: m for m in LegacyProductMap.objects.filter(tenant=self.a)}
        self.assertEqual((maps["T12A"].status, maps["T12A"].product_id), (MapStatus.CONFIRMED, self.used.pk))
        self.assertEqual((maps["T12OLD"].status, maps["T12OLD"].product_id), (MapStatus.UNMAPPED, None))
        self.assertEqual(maps["T4710000000001"].status, MapStatus.UNMAPPED)
        # 收回之後可以再建一次
        again = products.bring_in_products(self.a, [products.LegacyItem("T4710000000001", "皮套/B", 100, 1)],
                                           MAPPING, write=True)
        self.assertEqual(again.created, 1)

    def test_a_mapping_a_person_redid_is_not_taken_back(self):
        m = LegacyProductMap.objects.get(tenant=self.a, product_code_raw="T4710000000001")
        LegacyMappingLog.objects.create(tenant=self.a, kind=LegacyMappingLog.Kind.PRODUCT, source_system=SOURCE_SYSTEM,
                                        key_raw=m.product_code_raw, action=LegacyMappingLog.Action.CONFIRM,
                                        method=MapMethod.MANUAL, new_target_id=m.product_id, note="人確認")
        r = products.take_back_products(self.a, write=True)
        self.assertEqual(r, {"removed": 0, "unlinked": 1, "kept": 1})
        self.assertTrue(Product.objects.filter(pk=m.product_id).exists())


class CommandTests(TestCase):
    def setUp(self):
        self.a = Tenant.objects.create(name="甲通訊行", code="a")
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, True)
        self.report = self.tmp / "價格.xls"
        self.report.write_bytes(ored_page(HEAD_PRICE, price_rows(("T12A", "皮套/A", "350", "2"),
                                                                 ("AA01", "遠傳/月租", "0", "9999"))))
        self.cost = self.tmp / "存量.xls"
        self.cost.write_bytes(ored_page(HEAD_COST, [["T12A", "皮套/A", "1", "1", "120", "2", "240"]]))
        wb = Workbook()
        ws = wb.active
        ws.title = "類別對照"
        ws.append(["提議類別", "商品型態", "舊系統品名前綴", "庫存筆數", "你的修改(類別/排除)"])
        ws.append(["皮套", "配件", "皮套", 1, None])
        ws.append(["(排除)", "—", "遠傳", 1, "(排除)"])
        self.mapping = self.tmp / "對照.xlsx"
        wb.save(self.mapping)

    def call(self, *args):
        out = StringIO()
        call_command("bring_in_legacy_products", "--tenant", "a", *args, stdout=out)
        return out.getvalue()

    def test_trial_then_confirm_then_undo(self):
        files = ("--report", str(self.report), "--mapping", str(self.mapping))
        out = self.call(*files)
        self.assertIn("試算", out)
        self.assertIn("新建品號 1", out)
        self.assertIn("對照表填排除 1:遠傳 1", out)
        self.assertFalse(Product.objects.exists())
        self.assertIn("沒有成本的 1", out)
        self.assertIn("沒有成本的 0", self.call(*files, "--cost", str(self.cost), "--confirm"))
        p = Product.objects.get(barcode="T12A")
        self.assertEqual((p.name, p.weighted_avg_cost), ("皮套/A", Decimal("120")))
        self.assertIn("刪掉品號 1", self.call("--undo", "--confirm"))
        self.assertFalse(Product.objects.exists())

    def test_skip_list_leaves_those_old_codes_out(self):
        wb = Workbook()
        wb.active.title = "舊品號"
        drop = wb.create_sheet("未來會刪除")
        drop.append(["舊品號", "品名"])
        drop.append(["T12A", "皮套/A"])
        drop.append(["T99GONE", "皮套/早就沒了"])
        skip = self.tmp / "舊品號整理.xlsx"
        wb.save(skip)
        out = self.call("--report", str(self.report), "--mapping", str(self.mapping), "--skip", str(skip), "--confirm")
        self.assertIn("新建品號 0", out)
        self.assertIn("清單挑掉的 1(清單 2 個,這份報表沒有的 1)", out)
        self.assertFalse(Product.objects.exists())

    def test_needs_both_files_and_a_real_company(self):
        with self.assertRaises(CommandError):
            self.call()
        with self.assertRaises(CommandError):
            call_command("bring_in_legacy_products", "--tenant", "zz", "--undo", stdout=StringIO())
        with self.assertRaises(CommandError):
            self.call("--report", str(self.tmp / "沒有.xls"), "--mapping", str(self.mapping))
