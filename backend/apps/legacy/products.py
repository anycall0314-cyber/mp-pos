"""舊 POS 有庫存的品項整批建成品號(12/1 盤點入庫之前;owner 2026-10-14「照你建議」)。

- 只建品號:品名、條碼(= 舊品號,舊貼紙照刷)、售價(舊系統的零售價)、成本(舊系統「存量金額表」的成本,有給才帶)。
  **不帶數量** —— 盤點時刷舊貼紙打數量。成本放在商品的平均成本:盤點作業核准時,門市沒有過的商品就用它入庫
  (`inventory/stockcount.py`),所以配件的第一次入庫就是一張盤點單,不另外做入口。
- 類別照 owner 2026-10-05 填過的類別對照表(舊品名第一段 → 類別;填「排除」的不建、對照表沒有的列出來)。
- 手機、中古機不在這裡建:中古機一台一個舊品號(IMEI + I),新機在舊報表沒有逐支 IMEI;盤點時刷序號。
- 舊品號是「T + 13 碼」的,另外把那 13 碼記成這個商品的條碼別名(刷原廠條碼也找得到)。
- 已經對照過的舊品號不動;系統裡已經有這個條碼的商品 → 對照過去、不另外建;同名但條碼不同 → 不建、列出來。
- 每一筆都記在舊品號對照(`LegacyProductMap`,跟十年歷史共用同一張);收回只刪整批建的、還沒用過的。
- 有庫存 = 任何一家店數量大於 0(哪一家舊店對到哪一家新門市是盤點那一步的事,這裡不管)。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from html.parser import HTMLParser

from django.db import transaction
from django.db.models import ProtectedError
from django.utils import timezone

from .archive import SOURCE_SYSTEM
from .mapping import _log, _user
from .models import LegacyMappingLog, LegacyProductMap, MapMethod, MapStatus

# 對照表裡這幾個類別是一台一筆序號的,盤點時刷序號,不在這裡建
SERIAL_CATEGORIES = {"中古機", "手機"}
# 新建類別時用的代碼(品號前綴,只能英數);沒列到的用 G01、G02…
CATEGORY_CODES = {
    "背蓋": "BK", "玻璃貼": "PP", "保護貼": "SP", "空壓殼": "CA", "皮套": "LC", "鏡頭貼": "LS",
    "充電/線材": "CH", "耳機/音訊": "EP", "支架/車架": "ST", "掛繩/吊飾": "LN", "穿戴/周邊": "WB",
    "包膜": "FM", "維修料": "RP", "家電": "HA", "其他配件": "AC",
}
# 跟機型無關的類別:不用掛相容機型(不會出現在「待補相容機型」)
NO_HOST_MODEL = {"充電/線材", "耳機/音訊", "支架/車架", "掛繩/吊飾", "穿戴/周邊", "家電", "其他配件"}
# 維修用的零件:放零件倉(維修單只挑得到零件倉的商品)
PARTS_CATEGORIES = {"維修料"}
_GTIN_WITH_T = re.compile(r"T(\d{13})")
# 平均成本欄位放得下的(14 位、小數 2 位)
MAX_COST = Decimal("1e12")
# 對照紀錄的說明:收回只收這一句建的
BULK_NOTE = "整批建品號"


class ReportError(Exception):
    """看不懂這份報表(原因可以直接給人看)。"""


# ─────────────────────────── 讀報表 ───────────────────────────
class _Cells(HTMLParser):
    """歐睿匯出的「Excel」其實是網頁表格:一列一列的格子文字。"""

    def __init__(self):
        super().__init__()
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append("".join(self._cell).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _grid(raw: bytes) -> list[list[str]]:
    if raw[:2] == b"PK":   # 真的 xlsx
        from io import BytesIO

        from openpyxl import load_workbook

        ws = load_workbook(BytesIO(raw), read_only=True, data_only=True).active
        return [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("big5hkscs", "replace")
    parser = _Cells()
    parser.feed(text)
    return parser.rows


def read_ored_rows(raw: bytes) -> list[dict[str, str]]:
    """一份歐睿報表 → 一列一個 dict(欄名 → 格子文字)。表頭每隔幾十列會重複一次,跳過;序號不是數字的列(標題、小計)也跳過。"""
    header: list[str] | None = None
    out = []
    for cells in _grid(raw):
        if "品號" in cells and "品名" in cells:
            header = cells
            continue
        if header and cells and cells[0].strip().isdigit():
            out.append(dict(zip(header, cells)))
    if header is None:
        raise ReportError("找不到「品號」「品名」那一列表頭")
    return out


def _number(text) -> Decimal | None:
    try:
        return Decimal(str(text or "").replace(",", "").strip())
    except InvalidOperation:
        return None


@dataclass
class LegacyItem:
    code: str
    name: str
    price: int | None = None
    qty: int = 0
    cost: Decimal | None = None
    # 有沒有任何一家店數量大於 0(報表讀出來的才有;None = 照 qty 看)
    some_store: bool | None = None

    @property
    def in_stock(self) -> bool:
        return self.some_store if self.some_store is not None else self.qty > 0


# ST015 一列一個品號、一家店一欄:這幾欄以外的就是店
_NOT_STORES = {"序號", "類別", "分類", "品號", "品名", "共用", "零售價", "售價", "小計", "成本", "總數量", "總成本"}


def stock_items(rows: list[dict[str, str]]) -> list[LegacyItem]:
    """每個舊品號一筆:品名、零售價(四捨五入成整數;報表沒有這一欄就是 None)、各店數量加起來、有沒有任何一家店大於 0。

    兩種報表都認:庫存明細表(一列 = 一個品號 × 一家店,「數量」那一欄)、商品價格查詢表 ST015(一列一個品號,一家店一欄 +「小計」)。
    有庫存看的是「任何一家店大於 0」,不看加總(一家 +2、一家 −2 加起來是 0,東西還在)。
    """
    if rows and "數量" not in rows[0] and "小計" not in rows[0]:
        raise ReportError("這份報表沒有「數量」或「小計」")
    items: dict[str, LegacyItem] = {}
    for r in rows:
        code = (r.get("品號") or "").strip()
        if not code:
            continue
        item = items.setdefault(code, LegacyItem(code, (r.get("品名") or "").strip()))
        qty = _number(r.get("數量") if "數量" in r else r.get("小計")) or Decimal(0)
        item.qty += int(qty)
        if "數量" in r:
            here = [qty]
        else:
            here = [_number(v) or Decimal(0) for k, v in r.items() if k not in _NOT_STORES] or [qty]
        item.some_store = bool(item.some_store) or any(n > 0 for n in here)
        price = _number(r.get("零售價", r.get("售價")))
        if price is not None:
            price = int(price.quantize(Decimal(1), rounding=ROUND_HALF_UP))
            item.price = price if item.price is None else max(item.price, price)
    return list(items.values())


def old_costs(rows: list[dict[str, str]]) -> dict[str, Decimal]:
    """存量金額表:舊品號 → 一件的成本(照原數字,成本不收整數)。看不懂、負的不收。"""
    if rows and "成本" not in rows[0]:
        raise ReportError("這份報表沒有「成本」")
    out = {}
    for r in rows:
        code, cost = (r.get("品號") or "").strip(), _number(r.get("成本"))
        if code and cost is not None and 0 <= cost < MAX_COST:
            out[code] = cost.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)   # 平均成本存到分
    return out


def with_costs(items: list[LegacyItem], costs: dict[str, Decimal]) -> list[LegacyItem]:
    for item in items:
        item.cost = costs.get(item.code)
    return items


def prefix_of(name: str) -> str:
    return name.split("/", 1)[0].strip() if "/" in name else name.strip()


# ─────────────────────────── 建品號 ───────────────────────────
@dataclass
class Outcome:
    created: int = 0            # 新建的品號
    linked: int = 0             # 系統裡已經有這個條碼,對照過去
    already: int = 0            # 這個舊品號之前就對照好了
    no_price: int = 0           # 新建的裡面售價是 0 / 沒有售價的
    no_cost: int = 0            # 新建的裡面沒有成本(成本表沒有這個品號,或是 0)的
    aliases: int = 0            # 另外記了原廠條碼(T + 13 碼的那 13 碼)
    zero_stock: int = 0         # 沒有庫存,不建
    serial: int = 0             # 手機 / 中古機,盤點時刷序號
    excluded: dict = field(default_factory=dict)   # 對照表填「排除」的:{舊品名第一段: 幾個}
    unmapped: dict = field(default_factory=dict)   # 對照表沒有的
    same_name: list = field(default_factory=list)  # 同名但條碼不同,沒建:[(舊品號, 品名)]
    new_categories: list = field(default_factory=list)


def _category(tenant, name, cache, out):
    from apps.catalog.models import Category

    if name in cache:
        return cache[name]
    cat = Category.objects.filter(tenant=tenant, name=name).first()
    if cat is None:
        used = set(Category.objects.filter(tenant=tenant).values_list("code", flat=True))
        code = CATEGORY_CODES.get(name)
        if not code or code in used:
            n = 1
            while f"G{n:02d}" in used:
                n += 1
            code = f"G{n:02d}"
        cat = Category.objects.create(tenant=tenant, code=code, name=name,
                                      needs_host_model=name not in NO_HOST_MODEL)
        out.new_categories.append(f"{code} {name}")
    cache[name] = cat
    return cat


def _barcode_owner(tenant, barcode):
    """這個條碼現在指到哪一個商品(商品本身的條碼,或已確認的條碼別名);沒有就 None。"""
    from apps.catalog.models import Product
    from apps.identity.models import ProductAlias
    from apps.identity.normalize import alias_key

    pid = Product.objects.filter(tenant=tenant, barcode=barcode).values_list("pk", flat=True).first()
    if pid is None:
        pid = (ProductAlias.objects.filter(tenant=tenant, kind=ProductAlias.Kind.BARCODE, is_active=True,
                                           normalized_value=alias_key(barcode))
               .values_list("product_id", flat=True).first())
    return pid


def _confirm(tenant, item, product, method, user):
    row, _ = LegacyProductMap.objects.select_for_update().get_or_create(
        tenant=tenant, source_system=SOURCE_SYSTEM, product_code_raw=item.code,
        defaults={"product_name_seen": item.name[:500]},
    )
    row.product = product
    row.status = MapStatus.CONFIRMED
    row.method = method
    row.confirmed_by = _user(user)
    row.confirmed_at = timezone.now()
    if not row.product_name_seen:
        row.product_name_seen = item.name[:500]
    row.save()
    _log(tenant, LegacyMappingLog.Kind.PRODUCT, SOURCE_SYSTEM, item.code, LegacyMappingLog.Action.CONFIRM,
         method, None, product.pk, f"{product.sku} {product.name}", user, BULK_NOTE)


def bring_in_products(tenant, items: list[LegacyItem], mapping: dict[str, str], user=None, *, write=False) -> Outcome:
    """有庫存的舊品項建成品號。`mapping` = 類別對照表(舊品名第一段 → 類別,排除的不在裡面;見 `load_category_mapping`)。

    `write=False` 是試算:整段照做一次再全部退回,回的數字跟真的跑一樣。可以重跑(對照過的不動)。
    """
    from apps.catalog.models import Product
    from apps.identity.models import ProductAlias

    out = Outcome()
    excluded_prefixes = mapping.get("__excluded__", set())
    with transaction.atomic():
        done = set(LegacyProductMap.objects.filter(
            tenant=tenant, source_system=SOURCE_SYSTEM, status=MapStatus.CONFIRMED,
        ).values_list("product_code_raw", flat=True))
        names = set(Product.objects.filter(tenant=tenant).values_list("name", flat=True))
        cats: dict = {}
        for item in sorted(items, key=lambda i: i.code):
            if not item.in_stock:
                out.zero_stock += 1
                continue
            prefix = prefix_of(item.name)
            target = mapping.get(prefix)
            if target is None:
                bucket = out.excluded if prefix in excluded_prefixes else out.unmapped
                bucket[prefix] = bucket.get(prefix, 0) + 1
                continue
            if target in SERIAL_CATEGORIES:
                out.serial += 1
                continue
            if item.code in done:
                out.already += 1
                continue
            owner = _barcode_owner(tenant, item.code)
            if owner is not None:
                _confirm(tenant, item, Product.objects.get(pk=owner), MapMethod.CODE, user)
                out.linked += 1
                continue
            name = item.name or item.code
            if name in names:
                out.same_name.append((item.code, name))
                continue
            product = Product(
                tenant=tenant, category=_category(tenant, target, cats, out), name=name, barcode=item.code,
                list_price=item.price or 0, weighted_avg_cost=item.cost or 0,
                lifecycle_status=Product.LifecycleStatus.PENDING,
                requires_serial=False,   # 商品預設是逐台記序號;這裡都是按數量的配件
                warehouse_type=(Product.WarehouseType.PARTS if target in PARTS_CATEGORIES
                                else Product.WarehouseType.PRODUCT),
            )
            product.save()
            names.add(name)
            out.created += 1
            out.no_price += 0 if item.price else 1
            out.no_cost += 0 if item.cost else 1
            gtin = _GTIN_WITH_T.fullmatch(item.code)
            if gtin and _barcode_owner(tenant, gtin[1]) is None:
                ProductAlias.objects.create(
                    tenant=tenant, product=product, kind=ProductAlias.Kind.BARCODE, value=gtin[1],
                    source=ProductAlias.Source.IMPORTED, note=BULK_NOTE, created_by=_user(user),
                )
                out.aliases += 1
            _confirm(tenant, item, product, MapMethod.CREATED, user)
        if not write:
            transaction.set_rollback(True)
    return out


def take_back_products(tenant, user=None, *, write=False) -> dict:
    """把整批建的品號收回去:最新一筆對照紀錄是整批建的、現在還對照著那一個商品的才算。

    - 整批**新建**、還沒用過(沒有單、沒有庫存、沒有序號)→ 解除對照、刪掉那個品號;
    - 新建但已經用過、或刪不掉(被別的資料指著)→ 不動;
    - 當時是對照到**原本就有**的商品 → 只解除對照。
    回 {removed, unlinked, kept}。
    """
    from apps.catalog.models import Product
    from apps.catalog.usage import used_product_ids

    latest = {}
    for key, action, note, target in LegacyMappingLog.objects.filter(
        tenant=tenant, kind=LegacyMappingLog.Kind.PRODUCT, source_system=SOURCE_SYSTEM,
    ).order_by("pk").values_list("key_raw", "action", "note", "new_target_id"):
        latest[key] = (action, note, target)
    out = {"removed": 0, "unlinked": 0, "kept": 0}
    with transaction.atomic():
        rows = [
            m for m in LegacyProductMap.objects.select_for_update().filter(
                tenant=tenant, source_system=SOURCE_SYSTEM, status=MapStatus.CONFIRMED,
            ).order_by("pk")
            if latest.get(m.product_code_raw) == (LegacyMappingLog.Action.CONFIRM, BULK_NOTE, m.product_id)
        ]
        used = used_product_ids([m.product_id for m in rows if m.method == MapMethod.CREATED])
        for m in rows:
            created_here = m.method == MapMethod.CREATED
            if created_here and m.product_id in used:
                out["kept"] += 1
                continue
            old = m.product_id
            try:
                # 解除對照和刪品號一起:刪不掉(作廢的單、別的資料還指著它)就整段退回,對照也留著
                with transaction.atomic():
                    m.product, m.status, m.method, m.confirmed_by, m.confirmed_at = (
                        None, MapStatus.UNMAPPED, "", None, None)
                    m.save()
                    _log(tenant, LegacyMappingLog.Kind.PRODUCT, SOURCE_SYSTEM, m.product_code_raw,
                         LegacyMappingLog.Action.REVOKE, "", old, None, "", user, f"收回{BULK_NOTE}")
                    if created_here:
                        Product.objects.filter(pk=old).delete()
            except ProtectedError:
                out["kept"] += 1
                continue
            out["removed" if created_here else "unlinked"] += 1
        if not write:
            transaction.set_rollback(True)
    return out


def load_category_mapping(raw: bytes) -> dict:
    """owner 填過的類別對照表(xlsx,「類別對照」分頁):舊品名第一段 → 類別。
    「你的修改」那一欄有填就用它,沒填用「提議類別」;填排除的放進 `__excluded__`(試算時分開列)。"""
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    ws = wb["類別對照"] if "類別對照" in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows, ())]
    col = {h: i for i, h in enumerate(header)}
    try:
        i_prefix, i_proposed = col["舊系統品名前綴"], col["提議類別"]
    except KeyError:
        raise ReportError("類別對照表要有「舊系統品名前綴」「提議類別」兩欄") from None
    i_edit = next((i for h, i in col.items() if h.startswith("你的修改")), None)
    out: dict = {"__excluded__": set()}
    for row in rows:
        prefix = str(row[i_prefix] or "").strip() if i_prefix < len(row) else ""
        if not prefix:
            continue
        edit = str(row[i_edit] or "").strip() if i_edit is not None and i_edit < len(row) else ""
        final = edit or str(row[i_proposed] or "").strip()
        if final.strip("()()") in ("排除", "刪除", "不匯"):
            out["__excluded__"].add(prefix)
        elif final:
            out[prefix] = final
    return out
